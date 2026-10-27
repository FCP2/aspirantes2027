"""Genera el reporte documental de coordinaciones, sin modificar SQLite ni Excel.

Uso: python tools/reportes/coordinaciones_pdf.py [--inspect]
"""

from __future__ import annotations

import argparse
import collections
import io
import re
import sqlite3
import unicodedata
from urllib.parse import urlparse
from datetime import date, datetime
from html import escape
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image as PILImage, ImageOps
import pdfplumber
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import (
    Image,
    PageBreak,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "storage/reportes/registros_coordinaciones4tv1.xlsx"
PREVIOUS_SOURCE = ROOT / "storage/reportes/registros_coordinaciones4t.xlsx"
MATCHES = ROOT / "storage/reportes/coordinaciones4t_con_fotos.xlsx"
DATABASE = ROOT / "data/municipal.sqlite"
OUTPUT = ROOT / "output/pdf/coordinaciones4t_v1_por_municipio.pdf"
PHOTO_DIR = ROOT / "storage/photos"
FONT_DIR = Path("C:/Users/zaiko/AppData/Local/Microsoft/Windows/Fonts")
HEADER = (
    "INSCRIPCIONES REALIZADAS A LAS COORDINACIONES MUNICIPALES DE LOS "
    "COMITÉS DE DEFENSA DE LA TRANSFORMACIÓN Y LA SOBERANÍA NACIONAL "
    "| ESTADO DE MÉXICO"
)
BODY_TOP = 715
GUINDA = colors.HexColor("#86163F")
GUINDA_DARK = colors.HexColor("#65112F")
GOLD = colors.HexColor("#B79665")
BEIGE = colors.HexColor("#F3EEE6")
LINE = colors.HexColor("#DDD7D2")
INK = colors.HexColor("#25242A")
MUTED = colors.HexColor("#6F7075")


def normalized(value):
    value = unicodedata.normalize("NFKD", str(value or ""))
    value = "".join(char for char in value if not unicodedata.combining(char))
    return re.sub(r"[^A-Z0-9]+", " ", value.upper()).strip()


def clean(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def reference_field(value, label):
    match = re.search(rf"(?im)^{re.escape(label)}:\s*(.+)$", str(value or ""))
    return clean(match.group(1)) if match else ""


def reference_links(value):
    links = []
    seen = set()
    for found in re.findall(r"https?://[^\s<>\]\)]+", str(value or ""), flags=re.I):
        url = found.rstrip(".,;")
        parsed = urlparse(url)
        if parsed.scheme.lower() not in ("http", "https") or not parsed.netloc:
            continue
        if url not in seen:
            seen.add(url)
            links.append(url)
    return links


def link_label(url):
    parsed = urlparse(url)
    domain = parsed.netloc.removeprefix("www.")
    slug = parsed.path.strip("/").split("/")[-1]
    slug = re.sub(r"[-_]+", " ", slug).strip()
    if slug and len(slug) > 22:
        slug = slug[:21].rstrip() + "…"
    return f"{domain} - {slug}" if slug else domain


def date_text(value):
    if isinstance(value, (date, datetime)):
        months = (
            "ene", "feb", "mar", "abr", "may", "jun",
            "jul", "ago", "sep", "oct", "nov", "dic",
        )
        return f"{value.day:02d} {months[value.month - 1]} {value.year}"
    return clean(value) or "No indicado"


def photo_file(raw_path):
    if not raw_path:
        return None
    original = Path(raw_path)
    candidate = PHOTO_DIR / original.name
    return candidate if candidate.is_file() else None


def aspirant_index():
    connection = sqlite3.connect(f"file:{DATABASE.as_posix()}?mode=ro", uri=True)
    connection.row_factory = sqlite3.Row
    rows = connection.execute("""
        SELECT m.nombre AS municipio, p.nombre, p.apellido_paterno,
               p.apellido_materno, p.foto_path
        FROM actores_municipio a
        JOIN municipios m ON m.id = a.municipio_id
        JOIN personas p ON p.id = a.persona_id
        WHERE a.tipo_actor = 'ASPIRANTE_2027' AND a.activo = 1 AND p.activo = 1
    """).fetchall()
    connection.close()
    index = collections.defaultdict(list)
    for row in rows:
        name = " ".join(
            filter(None, (row["nombre"], row["apellido_paterno"], row["apellido_materno"]))
        )
        index[(normalized(row["municipio"]), normalized(name))].append(
            {"name": name, "photo": photo_file(row["foto_path"])}
        )
    return index


def source_identity(row):
    return (
        normalized(row[3]),
        normalized(row[4]),
        normalized(row[7]),
    )


def unique_photo(candidates):
    photos = {
        str(person["photo"]): person["photo"]
        for person in candidates if person["photo"]
    }
    return next(iter(photos.values())) if len(photos) == 1 else None


def photo_for_new_record(municipio, name, index):
    municipality_key = normalized(municipio)
    name_key = normalized(name)
    direct = unique_photo(index.get((municipality_key, name_key), []))
    if direct:
        return direct
    if len(name_key.split()) < 2:
        return None
    compatible = [
        person
        for (indexed_municipality, indexed_name), people in index.items()
        if indexed_municipality == municipality_key
        and f" {name_key} " in f" {indexed_name} "
        for person in people
    ]
    return unique_photo(compatible) if len(compatible) == 1 else None


def load_records():
    source = source_rows(SOURCE)
    previous = source_rows(PREVIOUS_SOURCE)
    matches = source_rows(MATCHES)
    if len(previous) != len(matches) or any(
        tuple(original[:12]) != tuple(result[:12])
        for original, result in zip(previous, matches)
    ):
        raise ValueError("El Excel del cruce no corresponde a la base anterior")
    if collections.Counter(map(source_identity, previous)) - collections.Counter(map(source_identity, source)):
        raise ValueError("La nueva base modificó o eliminó registros anteriores; revisar el cruce")

    index = aspirant_index()
    matches_by_identity = collections.defaultdict(collections.deque)
    for match in matches:
        matches_by_identity[source_identity(match)].append(match)

    result = []
    for excel_row, original in enumerate(source, start=2):
        previous_matches = matches_by_identity[source_identity(original)]
        match = previous_matches.popleft() if previous_matches else None
        status = clean(match[12]) if match else "Nuevo registro"
        photo = None
        display_name = clean(original[4])
        if status in ("Coincidencia exacta", "Variante compatible"):
            candidates = index.get((normalized(original[3]), normalized(match[13])), [])
            photo = unique_photo(candidates)
        elif not match:
            photo = photo_for_new_record(original[3], display_name, index)
        result.append({
            "row": excel_row,
            "municipio": clean(original[3]),
            "nombre": display_name,
            "sexo": clean(original[5]),
            "fecha": original[6],
            "folio": clean(original[7]),
            "entidad": clean(original[8]),
            "referente": clean(original[9]),
            "subreferente": clean(original[10]),
            "photo": photo,
            "status": status,
            "source_status": reference_field(original[11], "Estatus"),
            "profile": reference_field(original[11], "Cargo o perfil"),
            "links": reference_links(original[11]),
        })
    return result


def register_fonts():
    paths = {
        "Montserrat": FONT_DIR / "Montserrat-Regular.ttf",
        "Montserrat-Medium": FONT_DIR / "Montserrat-Medium.ttf",
        "Montserrat-SemiBold": FONT_DIR / "Montserrat-SemiBold.ttf",
        "Montserrat-Bold": FONT_DIR / "Montserrat-Bold.ttf",
    }
    for family, path in paths.items():
        if not path.is_file():
            raise FileNotFoundError(f"No se encuentra la fuente {path}")
        pdfmetrics.registerFont(TTFont(family, str(path)))
    pdfmetrics.registerFontFamily(
        "Montserrat",
        normal="Montserrat",
        bold="Montserrat-Bold",
    )


def styles():
    return {
        "eyebrow": ParagraphStyle(
            "Eyebrow", fontName="Montserrat-SemiBold", fontSize=8,
            leading=12, textColor=GUINDA, spaceAfter=8,
        ),
        "title": ParagraphStyle(
            "Title", fontName="Montserrat-Bold", fontSize=25,
            leading=31, textColor=INK, spaceAfter=16,
        ),
        "lead": ParagraphStyle(
            "Lead", fontName="Montserrat", fontSize=9.2,
            leading=15, textColor=INK, spaceAfter=18,
        ),
        "note": ParagraphStyle(
            "Note", fontName="Montserrat", fontSize=7.9,
            leading=12, textColor=MUTED, spaceAfter=8,
        ),
        "column": ParagraphStyle(
            "Column", fontName="Montserrat-SemiBold", fontSize=7.2,
            leading=10, textColor=colors.white,
        ),
        "name": ParagraphStyle(
            "Name", fontName="Montserrat-SemiBold", fontSize=8.8,
            leading=12.3, textColor=INK,
        ),
        "detail": ParagraphStyle(
            "Detail", fontName="Montserrat", fontSize=7.35,
            leading=10.5, textColor=INK,
        ),
        "no_photo": ParagraphStyle(
            "NoPhoto", fontName="Montserrat-Medium", fontSize=7,
            leading=10, textColor=MUTED, alignment=TA_CENTER,
        ),
    }


def photo_flowable(path, style):
    if not path:
        empty = Table([[Paragraph("SIN FOTO", style["no_photo"])]], colWidths=[60], rowHeights=[68])
        empty.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), BEIGE),
            ("BOX", (0, 0), (-1, -1), 0.4, LINE),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        return empty

    with PILImage.open(path) as original:
        photo = ImageOps.exif_transpose(original).convert("RGB")
    width, height = photo.size
    sample = photo.resize((min(width, 280), min(height, 320)))
    sw, sh = sample.size
    middle = list(range(round(sh * 0.15), max(round(sh * 0.85), 1)))
    dark = []
    for x in range(sw):
        pixels = [sample.getpixel((x, y)) for y in middle]
        dark.append(sum(max(pixel) < 34 for pixel in pixels) / max(len(pixels), 1) > 0.94)
    left = 0
    while left < round(sw * 0.28) and dark[left]:
        left += 1
    right = 0
    while right < round(sw * 0.28) and dark[sw - right - 1]:
        right += 1
    if left >= round(sw * 0.025) or right >= round(sw * 0.025):
        x0 = round(width * left / sw)
        x1 = width - round(width * right / sw)
        if x1 - x0 >= width * 0.45:
            photo = photo.crop((x0, 0, x1, height))

    # Encajar, no recortar el rostro; el fondo beige evita las franjas negras.
    fitted = ImageOps.contain(photo, (120, 136), PILImage.Resampling.LANCZOS)
    canvas = PILImage.new("RGB", (120, 136), (247, 244, 240))
    canvas.paste(fitted, ((120 - fitted.width) // 2, (136 - fitted.height) // 2))
    buffer = io.BytesIO()
    canvas.save(buffer, format="JPEG", quality=88, optimize=True)
    buffer.seek(0)
    return Image(buffer, width=60, height=68)


def paragraph(value, style):
    return Paragraph(escape(str(value)), style)


def row_details(record, style):
    lines = []
    if record["profile"]:
        lines.append(("CARGO / PERFIL", record["profile"]))
    if record["sexo"]:
        lines.append(("SEXO / EXPRESIÓN", record["sexo"]))
    if record["source_status"]:
        lines.append(("ESTATUS", record["source_status"]))
    lines.append(("REGISTRO", date_text(record["fecha"])))
    lines.append(("FOLIO", record["folio"] or "No indicado"))
    if record["referente"]:
        lines.append(("REFERENTE", record["referente"]))
    if record["subreferente"]:
        lines.append(("SUBREFERENTE", record["subreferente"]))
    lines.append(("FUENTES", f"Excel de registros, fila {record['row']}"))
    html = "<br/>".join(
        f'<font color="#86163F"><b>{escape(label)}:</b></font>&#160;{escape(value)}'
        for label, value in lines
    )
    for number, url in enumerate(record["links"], start=1):
        label = f"{number}. {link_label(url)}"
        html += (
            '<br/><font color="#86163F">'
            f'<link href="{escape(url, quote=True)}"><u>{escape(label)}</u></link>'
            "</font>"
        )
    return Paragraph(html, style["detail"])


def draw_chrome(canvas, page_number):
    canvas.saveState()
    width, height = letter
    header_style = ParagraphStyle(
        "PageHeader", fontName="Montserrat-SemiBold", fontSize=7.6,
        leading=10.5, textColor=GUINDA,
    )
    header = Paragraph(escape(HEADER), header_style)
    _, header_height = header.wrap(522, 40)
    if header_height > 32:
        raise AssertionError("El encabezado ocupa más de tres líneas")
    header.drawOn(canvas, 45, height - 20 - header_height)
    canvas.setStrokeColor(GOLD)
    canvas.setLineWidth(1.2)
    canvas.line(45, 734, width - 45, 734)
    canvas.setFont("Montserrat-Medium", 8)
    canvas.setFillColor(GUINDA)
    canvas.drawRightString(width - 45, 29, f"{page_number}")
    canvas.restoreState()


def cover(records, grouped, style):
    photo_count = sum(bool(record["photo"]) for record in records)
    named_count = sum(bool(record["nombre"]) for record in records)
    empty_count = len(records) - named_count
    story = [
        Spacer(1, 36),
        paragraph("RELACIÓN DOCUMENTAL", style["eyebrow"]),
        paragraph("Registros para coordinaciones municipales", style["title"]),
        paragraph(
            "Personas mencionadas en la relación recibida, organizadas por municipio. "
            "Las fotografías se incorporaron únicamente cuando hubo una coincidencia "
            "de identidad verificable con el bloque Aspirantes 2027 de las fichas municipales.",
            style["lead"],
        ),
        Spacer(1, 10),
    ]
    metrics = [
        [paragraph("CONCEPTO", style["column"]), paragraph("TOTAL", style["column"])],
        [paragraph("Personas con nombre en la relación", style["detail"]), paragraph(named_count, style["name"])],
        [paragraph("Municipios representados", style["detail"]), paragraph(len(grouped), style["name"])],
        [paragraph("Fotografías vinculadas", style["detail"]), paragraph(photo_count, style["name"])],
        [paragraph("Municipios sin nombre aportado", style["detail"]), paragraph(empty_count, style["name"])],
    ]
    table = Table(metrics, colWidths=[375, 147], hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), GUINDA),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, BEIGE]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 12),
        ("TOPPADDING", (0, 0), (-1, -1), 11),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 11),
        ("LINEBELOW", (0, 1), (-1, -2), 0.35, LINE),
        ("BOX", (0, 0), (-1, -1), 0.6, LINE),
    ]))
    story.extend([
        table,
        Spacer(1, 22),
        paragraph(
            "El estatus de cada persona reproduce la clasificación del Excel fuente. "
            "Cuando no se encontró fotografía o el cruce requería revisión, se dejó "
            "el espacio sin imagen. La documentación completa de cada registro "
            "permanece en el Excel de origen; cada renglón de este reporte indica su fila "
            "y muestra los vínculos de corroboración disponibles.",
            style["note"],
        ),
        PageBreak(),
    ])
    return story


def draw_cover(canvas, records, grouped, style):
    y = BODY_TOP
    for flowable in cover(records, grouped, style):
        if isinstance(flowable, PageBreak):
            break
        _, height = flowable.wrap(522, y - 56)
        y -= height
        flowable.drawOn(canvas, 45, y)
        y -= flowable.getSpaceAfter()
    if y < 56:
        raise AssertionError("La portada invade el margen inferior")


def draw_municipality_header(canvas, number, name, count, top, continued=False):
    x = 45
    width = 522
    title_height = 26
    heading_height = 24
    title = f"{number:02d} · {name.upper()}"
    if continued:
        title += " · CONTINÚA"

    canvas.saveState()
    canvas.setFillColor(BEIGE)
    canvas.rect(x, top - title_height, width, title_height, fill=1, stroke=0)
    canvas.setStrokeColor(LINE)
    canvas.setLineWidth(0.5)
    canvas.rect(x, top - title_height, width, title_height, fill=0, stroke=1)
    font_size = 10.2
    while pdfmetrics.stringWidth(title, "Montserrat-Bold", font_size) > 395:
        font_size -= 0.2
    canvas.setFont("Montserrat-Bold", font_size)
    canvas.setFillColor(GUINDA)
    canvas.drawString(x + 10, top - 17, title)
    canvas.setFont("Montserrat-Medium", 8)
    canvas.drawRightString(x + width - 10, top - 16, f"{count} PERSONAS")

    columns_top = top - title_height
    canvas.setFillColor(GUINDA)
    canvas.rect(x, columns_top - heading_height, width, heading_height, fill=1, stroke=0)
    canvas.setFillColor(colors.white)
    canvas.setFont("Montserrat-SemiBold", 7.2)
    canvas.drawString(x + 9, columns_top - 15, "FOTO")
    canvas.drawString(x + 76 + 9, columns_top - 15, "PARTICIPANTE")
    canvas.drawString(x + 76 + 159 + 9, columns_top - 15, "REGISTRO Y VÍNCULOS")
    canvas.restoreState()
    return columns_top - heading_height


def record_layout(record, style):
    if not record["nombre"]:
        return {"height": 32, "photo": None, "name": None, "details": None}
    photo = photo_flowable(record["photo"], style)
    photo.wrap(60, 68)
    name = paragraph(record["nombre"], style["name"])
    details = row_details(record, style)
    _, name_height = name.wrap(143, 1000)
    _, details_height = details.wrap(271, 1000)
    height = max(68, name_height, details_height) + 18
    if height > 620:
        raise ValueError(f"El registro de la fila {record['row']} no cabe en una página")
    return {
        "height": height,
        "photo": photo,
        "name": name,
        "name_height": name_height,
        "details": details,
        "details_height": details_height,
    }


def draw_record(canvas, record, layout, top, alternate=False):
    x = 45
    width = 522
    height = layout["height"]
    bottom = top - height
    if bottom < 56:
        raise AssertionError(f"La fila {record['row']} invade el pie de página")

    canvas.saveState()
    canvas.setFillColor(colors.HexColor("#FCFAF8") if alternate else colors.white)
    canvas.rect(x, bottom, width, height, fill=1, stroke=0)
    canvas.setStrokeColor(LINE)
    canvas.setLineWidth(0.4)
    canvas.rect(x, bottom, width, height, fill=0, stroke=1)
    if not record["nombre"]:
        canvas.setFont("Montserrat", 7.35)
        canvas.setFillColor(INK)
        canvas.drawString(x + 8, bottom + 11, f"Sin persona identificada en el Excel base. Fila {record['row']}.")
        canvas.restoreState()
        return bottom

    # Cada celda dispone de su propia zona de dibujo; ningún elemento puede
    # invadir visualmente el renglón siguiente aun con visores PDF distintos.
    cell_areas = (
        (x, 76),
        (x + 76, 159),
        (x + 76 + 159, 287),
    )
    positions = (
        (layout["photo"], x + 8, bottom + (height - 68) / 2),
        (layout["name"], x + 76 + 8, bottom + (height - layout["name_height"]) / 2),
        (layout["details"], x + 76 + 159 + 8, top - 9 - layout["details_height"]),
    )
    for (cell_x, cell_width), (flowable, draw_x, draw_y) in zip(cell_areas, positions):
        canvas.saveState()
        clip = canvas.beginPath()
        clip.rect(cell_x + 1, bottom + 1, cell_width - 2, height - 2)
        canvas.clipPath(clip, stroke=0, fill=0)
        flowable.drawOn(canvas, draw_x, draw_y)
        canvas.restoreState()
    canvas.restoreState()
    return bottom


def build():
    register_fonts()
    records = load_records()
    grouped = collections.OrderedDict()
    for record in records:
        grouped.setdefault(record["municipio"], []).append(record)
    style = styles()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    canvas = Canvas(
        str(OUTPUT), pagesize=letter, pageCompression=1,
        invariant=1,
    )
    canvas.setTitle("Registros de coordinaciones municipales por municipio")
    canvas.setAuthor("Sistema de Fichas Municipales")
    canvas.setSubject("Relación documental con fotografías vinculadas a Aspirantes 2027")
    page_number = 1
    draw_chrome(canvas, page_number)
    draw_cover(canvas, records, grouped, style)
    canvas.showPage()
    page_number += 1
    draw_chrome(canvas, page_number)
    y = BODY_TOP
    rendered_rows = []

    for number, (municipio, people) in enumerate(grouped.items(), start=1):
        layouts = [(record, record_layout(record, style)) for record in people]
        count = sum(bool(person["nombre"]) for person in people)
        first_needed = 50 + layouts[0][1]["height"] + 14
        group_needed = 50 + sum(layout["height"] for _, layout in layouts) + 14
        needed = group_needed if count <= 2 and group_needed <= 660 else first_needed
        if y - needed < 56:
            canvas.showPage()
            page_number += 1
            draw_chrome(canvas, page_number)
            y = BODY_TOP
        y = draw_municipality_header(canvas, number, municipio, count, y)
        for index, (record, layout) in enumerate(layouts):
            if y - layout["height"] < 56:
                canvas.showPage()
                page_number += 1
                draw_chrome(canvas, page_number)
                y = draw_municipality_header(
                    canvas, number, municipio, count, BODY_TOP, continued=True
                )
            top = y
            y = draw_record(canvas, record, layout, y, alternate=index % 2 == 1)
            rendered_rows.append((record["row"], page_number, top, y))
        y -= 14

    for current, following in zip(rendered_rows, rendered_rows[1:]):
        if current[1] == following[1] and following[2] > current[3] + 0.01:
            raise AssertionError(f"Renglones superpuestos: {current[0]} y {following[0]}")
    if len(rendered_rows) != len(records):
        raise AssertionError("No se dibujaron todas las filas del Excel")
    canvas.save()
    print("output", OUTPUT)
    print("pages", page_number)
    print("source_rows", len(records), "named_people", sum(bool(record["nombre"]) for record in records), "municipalities", len(grouped))
    print("photos", sum(bool(record["photo"]) for record in records))
    print("source_links", sum(len(record["links"]) for record in records))


def verify():
    records = load_records()
    reader = PdfReader(str(OUTPUT))
    texts = [page.extract_text() or "" for page in reader.pages]
    whole = "\n".join(texts)
    for number, text in enumerate(texts, start=1):
        if "INSCRIPCIONES REALIZADAS A LAS COORDINACIONES MUNICIPALES" not in text:
            raise AssertionError(f"Falta el nuevo encabezado en la página {number}")
        if "REGISTROS DE COORDINACIONES MUNICIPALES" in text:
            raise AssertionError(f"Permanece el encabezado anterior en la página {number}")
        if "Fuente: Excel de registros de coordinaciones" in text:
            raise AssertionError(f"Permanece la fuente en el pie de la página {number}")
    missing = []
    for record in records:
        marker = (
            f"Excel de registros, fila {record['row']}"
            if record["nombre"]
            else f"Fila {record['row']}."
        )
        if marker not in whole:
            missing.append(record["row"])
    if missing:
        raise AssertionError(f"Faltan filas del Excel: {missing[:20]}")
    if any(len(text.strip()) < 100 for text in texts):
        raise AssertionError("Hay páginas en blanco o con muy poco contenido")
    if any(
        abs(float(page.mediabox.width) - 612) > 0.1
        or abs(float(page.mediabox.height) - 792) > 0.1
        for page in reader.pages
    ):
        raise AssertionError("El tamaño de página no es carta")
    expected_links = collections.Counter(
        url for record in records for url in record["links"]
    )
    actual_links = collections.Counter()
    link_pages = collections.Counter()
    for number, page in enumerate(reader.pages, start=1):
        for annotation_ref in page.get("/Annots", []):
            action = annotation_ref.get_object().get("/A")
            if action and action.get("/URI"):
                actual_links[str(action["/URI"])] += 1
                link_pages[number] += 1
    if set(actual_links) != set(expected_links) or any(
        actual_links[url] < count for url, count in expected_links.items()
    ):
        raise AssertionError(
            f"Enlaces incompletos: esperados {sum(expected_links.values())}, "
            f"presentes {sum(actual_links.values())}"
        )
    footer_intrusions = []
    image_count = 0
    image_overlap_pages = []
    image_outside_cells = []
    with pdfplumber.open(str(OUTPUT)) as pdf:
        for number, page in enumerate(pdf.pages, start=1):
            if f"\n{number}\n" not in f"\n{texts[number - 1]}\n":
                raise AssertionError(f"Falta la numeración de la página {number}")
            footer_text = page.crop((0, 745, 612, 792)).extract_text().strip()
            if footer_text != str(number):
                raise AssertionError(f"El pie de la página {number} contiene más que el numeral: {footer_text!r}")
            intruding = [char for char in page.chars if 738 < char["top"] < 750]
            if intruding:
                footer_intrusions.append(number)
            photos = sorted(page.images, key=lambda image: image["top"])
            image_count += len(photos)
            for photo in photos:
                if (
                    photo["x0"] < 45 or photo["x1"] > 121
                    or photo["top"] < 60 or photo["bottom"] > 736
                ):
                    image_outside_cells.append(number)
            for current, following in zip(photos, photos[1:]):
                if current["bottom"] > following["top"] + 0.5:
                    image_overlap_pages.append(number)
    if footer_intrusions:
        raise AssertionError(f"Contenido invadiendo el pie: {footer_intrusions[:20]}")
    if image_count != sum(bool(record["photo"]) for record in records):
        raise AssertionError(f"Fotografías incompletas: {image_count}")
    if image_overlap_pages or image_outside_cells:
        raise AssertionError(
            f"Fotografías superpuestas: {image_overlap_pages[:20]}; "
            f"fuera de su columna: {image_outside_cells[:20]}"
        )
    print("verified_pages", len(reader.pages))
    print("verified_excel_rows", len(records))
    print("verified_source_links", sum(expected_links.values()))
    print("clickable_areas", sum(actual_links.values()))
    print("pages_with_most_links", link_pages.most_common(5))
    print("footer_intrusions", len(footer_intrusions))
    print("verified_photos", image_count)
    print("photo_overlaps", len(image_overlap_pages))


def source_rows(path: Path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    workbook.close()
    return rows


def inspect():
    rows = source_rows(SOURCE)
    previous = source_rows(PREVIOUS_SOURCE)
    matched = source_rows(MATCHES)
    same = len(previous) == len(matched) and all(
        tuple(original[:12]) == tuple(result[:12])
        for original, result in zip(previous, matched)
    )
    new_people = collections.Counter(map(source_identity, rows)) - collections.Counter(map(source_identity, previous))
    print("source_rows", len(rows), "previous_rows", len(previous), "matched_rows", len(matched), "previous_match_valid", same)
    print("added_rows", sum(new_people.values()))
    print("municipalities", len({str(row[3]) for row in rows}))
    print("match_statuses", collections.Counter(str(row[12]) for row in matched))
    print("references_max_length", max(len(str(row[11] or "")) for row in rows))
    print("source_status_count", sum(bool(re.search(r"(?im)^estatus:\s*", str(row[11] or ""))) for row in rows))
    print("profile_count", sum(bool(re.search(r"(?im)^cargo o perfil:\s*", str(row[11] or ""))) for row in rows))
    print("source_statuses", collections.Counter(clean(re.search(r"(?im)^estatus:\s*(.+)$", str(row[11] or "")).group(1)) for row in rows).most_common(10))
    urls = []
    for index, row in enumerate(rows, start=2):
        for url in re.findall(r"https?://[^\s<>\]\)]+", str(row[11] or "")):
            urls.append((index, url.rstrip(".,;")))
    print("url_count", len(urls), "url_rows", len({row for row, _ in urls}))
    print("url_domains", collections.Counter(urlparse(url).netloc for _, url in urls).most_common(12))
    print("url_examples", urls[:3])
    by_row = collections.Counter(row for row, _ in urls)
    print("max_urls_per_row", by_row.most_common(8))
    print("longest_url_length", max((len(url) for _, url in urls), default=0))
    workbook = load_workbook(SOURCE, read_only=False, data_only=False)
    sheet = workbook.active
    linked = [(cell.coordinate, cell.hyperlink.target) for row in sheet for cell in row if cell.hyperlink and cell.hyperlink.target]
    print("excel_hyperlink_count", len(linked), "examples", linked[:3])
    workbook.close()
    print("references_sample", str(rows[0][11])[:700])
    print("sample_records")
    for index in (0, 1, 2, 10, 50, 150, 300, 500):
        row = rows[index]
        print(index + 2, [str(value)[:90] for value in row[3:11]])
    print("source_columns", rows[0][:12])
    print("max_name", max(len(str(row[4] or "")) for row in rows))
    print("max_referente", max(len(str(row[9] or "")) for row in rows))
    print("empty_names", [(index, row[3], row[7]) for index, row in enumerate(rows, start=2) if not clean(row[4])])
    print("empty_name_municipality_rows", {row[3]: sum(other[3] == row[3] for other in rows) for row in rows if not clean(row[4])})
    for column, label in ((3, "municipio"), (4, "nombre"), (9, "referente"), (10, "subreferente")):
        print(label, "replacement_count", sum("�" in str(row[column] or "") for row in rows))
    connection = sqlite3.connect(f"file:{DATABASE.as_posix()}?mode=ro", uri=True)
    for table in ("municipios", "personas", "actores_municipio"):
        print(table, connection.execute(f"PRAGMA table_info({table})").fetchall())
    print(
        "aspirant_types",
        connection.execute(
            "SELECT tipo_actor, COUNT(*) FROM actores_municipio GROUP BY tipo_actor"
        ).fetchall(),
    )
    print(
        "photo_examples",
        connection.execute(
            "SELECT foto_path FROM personas WHERE foto_path IS NOT NULL LIMIT 5"
        ).fetchall(),
    )
    connection.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--inspect", action="store_true")
    parser.add_argument("--verify", action="store_true")
    args = parser.parse_args()
    if args.inspect:
        inspect()
        return
    if args.verify:
        verify()
        return
    build()


if __name__ == "__main__":
    main()
